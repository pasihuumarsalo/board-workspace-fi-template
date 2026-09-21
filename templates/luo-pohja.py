# -*- coding: utf-8 -*-
"""Luo puhtaan pöytäkirjapohjan poytakirjapohja.docx.

    python templates/luo-pohja.py

Ajetaan kerran; sen jälkeen pohjaa muokataan Wordissa (yhtiön nimi, paikka,
osallistujat, allekirjoituslauseke). Rakennetta ei saa muuttaa: numeroitu lista,
[ASIA #]-merkinnät ja [HH:MM]/[DD.MM.YYYY]-paikkamerkit ovat generaattorin
etsimiä kohtia.
"""
import os
import docx
from docx.shared import Pt, Cm

JUURI = os.path.dirname(os.path.abspath(__file__))
ULOS = os.path.join(JUURI, 'poytakirjapohja.docx')

d = docx.Document()
for s in ('Normal', 'List Number'):
    d.styles[s].font.name = 'Calibri'
    d.styles[s].font.size = Pt(11)

# Ylätunniste
d.sections[0].header.paragraphs[0].text = '[Yhtiön nimi]  ·  Hallituksen kokouksen pöytäkirja'

def rivi(otsikko, arvo, rooli=None):
    p = d.add_paragraph()
    p.add_run(otsikko)
    p.add_run().add_tab(); p.add_run().add_tab()
    p.add_run(arvo)
    if rooli:
        p.add_run().add_tab(); p.add_run(rooli)
    return p

rivi('Paikka', '[Kokouspaikka]')
rivi('Päivä', '[DD.MM.YYYY]')
rivi('Osallistujat', '[Etunimi Sukunimi]', 'puheenjohtaja')
for _ in range(3):
    p = d.add_paragraph()
    p.add_run().add_tab(); p.add_run().add_tab()
    p.add_run('[Etunimi Sukunimi]'); p.add_run().add_tab(); p.add_run('jäsen')

d.add_heading('Kokous numero [#]', level=1)

def pykala(teksti):
    return d.add_paragraph(teksti, style='List Number')

pykala('Puheenjohtaja avasi kokouksen klo [HH:MM].')
pykala('Todettiin, että kokous oli kutsuttu koolle yhtiöjärjestyksen mukaisesti ja että '
       'paikalla oli [#] jäsentä. Kokous todettiin lailliseksi ja päätösvaltaiseksi.')
pykala('Hyväksyttiin kokouksen esityslista työjärjestykseksi.')
pykala('[ASIA #]'); pykala('[ASIA #]'); pykala('[ASIA #]')
pykala('Puheenjohtaja päätti kokouksen klo [HH:MM].')

d.add_paragraph()
rivi('Vakuudeksi', 'Tämä pöytäkirja on allekirjoitettu sähköisesti ja sisältää erillisen '
     'allekirjoitussivun, jossa allekirjoituksen varmenne on näkyvissä.')

d.save(ULOS)
print('luotu:', ULOS)
