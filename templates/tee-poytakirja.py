# -*- coding: utf-8 -*-
"""Muodostaa hallituksen pöytäkirjan — tai päätöksen ilman kokousta — päätösrekisteristä
Word-pohjalle.

    python templates/tee-poytakirja.py PTK-000073 [ulostulo.docx] [--valmis] [--rekisteri polku.yaml]

Muokkaa pohjaa PAIKOILLEEN kopiosta: tyylit ja Wordin automaattinen pykälänumerointi
periytyvät. Ulkoasumuutos tehdään siis Word-pohjaan, ei tähän. Kohdat etsitään
sisällön perusteella, ei kappalenumeron.

Kaikki henkilötiedot tulevat REKISTERISTÄ: osallistujat, sijaiset, poissaolijat,
toimitusjohtaja, sihteeri, kutsutut, esteellisyydet, äänestykset, eriävät mielipiteet
ja allekirjoittajat. Pohjan paikkamerkkinimet korvataan, ei koskaan käytetä.

Ennen asiakirjaa ajetaan sama validointi kuin scripts/validoi-tyotila.py:ssä.
  * Oletus on LUONNOS: syntyy aina, banneri kertoo virheiden ja varoitusten määrän.
  * --valmis EI ole pelkkä bannerin poisto. Se edellyttää, ettei yhtään ERROR-tason
    löydöstä ole — mukaan lukien päätösvaltaisuus jokaisessa asiassa, esteellisyys
    kirjattuna, puheenjohtajan sanamuotovahvistus, allekirjoittajat ja koollekutsu.
    Jos yksikin estävä tarkistus epäonnistuu, generaattori kieltäytyy (paluuarvo 2),
    tulostaa virheluettelon eikä tuota mitään. Tätä ei voi ohittaa parametrilla.
"""
import copy
import os
import re
import shutil
import sys

import docx
from docx.shared import Cm
from docx.text.paragraph import Paragraph

JUURI = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(JUURI, 'scripts'))
import tyotila as T  # noqa: E402

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
POHJAT = {'meeting': os.path.join(JUURI, 'templates', 'poytakirjapohja.docx'),
          'without_meeting': os.path.join(JUURI, 'templates', 'paatospohja-ilman-kokousta.docx')}
NIMIPAIKKA = '[Etunimi Sukunimi]'
ROOLINIMET = {'chair': 'puheenjohtaja', 'vice_chair': 'varapuheenjohtaja', 'member': 'jäsen', 'deputy': 'varajäsen'}


class Kieltaytyi(SystemExit):
    """Valmista asiakirjaa ei muodostettu: validointi löysi estäviä virheitä."""


# ------------------------------------------------------------------ Word-apurit

def juokseva_numero(ptk):
    """PTK-000073 -> '73'. Asiakirjalla numero on selkokielinen."""
    return str(T.tunnus_osat(ptk)[1])


def korvaa(p, vanha, uusi):
    """Korvaa merkkijonon kappaleessa vaikka se olisi jakautunut useaan runiin."""
    runs = p.runs
    koko = ''.join(r.text for r in runs)
    if vanha not in koko:
        return False
    a = koko.index(vanha); b = a + len(vanha); pos = 0
    for r in runs:
        rt = r.text; rs, re_ = pos, pos + len(rt); pos = re_
        if re_ <= a or rs >= b:
            continue
        ls = max(a, rs) - rs; le = min(b, re_) - rs
        r.text = (rt[:ls] + uusi + rt[le:]) if rs <= a else (rt[:ls] + rt[le:])
    return True


def etsi(d, hakusana, pakollinen=True):
    for p in d.paragraphs:
        if hakusana in p.text:
            return p
    if pakollinen:
        raise SystemExit('Pohjasta ei löytynyt kohtaa: %r' % hakusana)
    return None


def poista(p):
    p._p.getparent().remove(p._p)


def kloonaa_jalkeen(malli, ankkuri):
    el = copy.deepcopy(malli._p)
    ankkuri._p.addnext(el)
    return Paragraph(el, malli._parent)


def kloonaa_asia(malli, teksti, lihava=False, numeroitu=True):
    """Kloonaa asiapykälän mallin. Numeroidut perivät listatyylin; muut saavat Normal-tyylin ja sisennyksen."""
    el = copy.deepcopy(malli._p)
    for r in el.findall(W + 'r'):
        el.remove(r)
    p = Paragraph(el, malli._parent)
    if not numeroitu:
        pPr = el.find(W + 'pPr')
        if pPr is not None:
            npr = pPr.find(W + 'numPr')
            if npr is not None:
                pPr.remove(npr)
        p.style = malli.part.document.styles['Normal']
        p.paragraph_format.left_indent = malli.paragraph_format.left_indent or Cm(1.27)
        p.paragraph_format.first_line_indent = Cm(0)
    p.add_run(teksti).bold = lihava
    return p


def jaa_kohdiksi(teksti):
    """Normalisoi välilyönnit ja jakaa '1. ' -luettelon omiksi kappaleikseen. Ei muuta sanamuotoa."""
    t = re.sub(r'\s+', ' ', teksti or '').strip()
    osat = re.split(r'(?=(?<![\d.])\d{1,2}\.\s+[A-ZÅÄÖ])', t)
    return [o.strip() for o in osat if o.strip()]


def taydenna_nimirivit(d, otsikko, rivit, paikka=NIMIPAIKKA, roolipaikat=('puheenjohtaja', 'jäsen', '[rooli]')):
    """Täyttää monistuvan nimi+rooli-lohkon: ensimmäinen rivi on otsikkorivi, sitä seuraavat
    paikkamerkkirivit monistetaan tai poistetaan tarpeen mukaan. Palauttaa False jos pohjassa
    ei ole lohkoa."""
    alku = None
    for p in d.paragraphs:
        if otsikko in p.text and paikka in p.text:
            alku = p
            break
    if alku is None:
        return False
    kaikki = d.paragraphs
    i = [q._p for q in kaikki].index(alku._p)
    paikat = [alku]
    for q in kaikki[i + 1:]:
        if paikka in q.text and otsikko not in q.text:
            paikat.append(q)
        else:
            break
    malli = paikat[1] if len(paikat) > 1 else paikat[0]
    tayta = [alku]
    ankkuri = paikat[-1]
    for _ in rivit[1:]:
        uusi = kloonaa_jalkeen(malli, ankkuri)
        ankkuri = uusi
        tayta.append(uusi)
    for p, (nimi_, rooli) in zip(tayta, rivit):
        korvaa(p, paikka, nimi_)
        for rp in roolipaikat:
            if korvaa(p, rp, rooli):
                break
    if not rivit:
        korvaa(alku, paikka, '—')
        for rp in roolipaikat:
            korvaa(alku, rp, '')
    for p in paikat[1:]:
        poista(p)
    return True


# ------------------------------------------------------------------ sisältö rekisteristä

def osallistujarivit(ty, m):
    k = ty.get('kokoonpano') or {}
    paiva = m.get('decision_date')
    roolit = m.get('roles') or {}
    pj = roolit.get('meeting_chair_person_id')
    las = T.lasnaolijat(m)
    tj = T.toimitusjohtaja_paivana(k, paiva)
    siht = roolit.get('secretary_person_id')
    rivit = []

    def lisarooli(pid, r):
        if pid == tj:
            r += ', toimitusjohtaja'
        if pid == siht:
            r += ', sihteeri'
        return r

    if pj:
        rk = T.rooli_paivana(k, pj, paiva)
        r = 'puheenjohtaja' if rk in ('chair', None) else 'puheenjohtaja (%s)' % ROOLINIMET.get(rk, rk)
        rivit.append((T.nimi(k, pj), lisarooli(pj, r)))
    for pid in (m.get('participants') or []):
        if pid == pj:
            continue
        rk = T.rooli_paivana(k, pid, paiva)
        rivit.append((T.nimi(k, pid), lisarooli(pid, ROOLINIMET.get(rk, 'jäsen') if rk != 'chair' else 'jäsen (hallituksen puheenjohtaja)')))
    for s in (m.get('deputy_substitutions') or []):
        if s.get('deputy_person_id') == pj:
            continue
        rivit.append((T.nimi(k, s['deputy_person_id']), 'varajäsen, sijaisena (%s)' % T.nimi(k, s.get('for_person_id'))))
    osall = roolit.get('ceo_participation')
    if tj and tj not in las and osall in ('present', 'partial'):
        rivit.append((T.nimi(k, tj), 'toimitusjohtaja' + (', osittain' if osall == 'partial' else '')))
    if siht and siht not in las and siht != tj:
        rivit.append((T.nimi(k, siht), 'sihteeri'))
    for inv in (m.get('invited_participants') or []):
        if isinstance(inv, str):
            rivit.append((inv, 'kutsuttu'))
        else:
            n = inv.get('display_name') or (T.nimi(k, inv['person_id']) if inv.get('person_id') else None) or inv.get('role') or 'kutsuttu'
            r = inv.get('role') or 'kutsuttu'
            if inv.get('items'):
                r += ' (%s)' % inv['items']
            rivit.append((n, r if n != r else 'kutsuttu'))
    poissa = [T.nimi(k, p) for p in T.poissaolijat(m)]
    return rivit, poissa


def allekirjoittajarivit(ty, m):
    k = ty.get('kokoonpano') or {}
    pj = (m.get('roles') or {}).get('meeting_chair_person_id')
    rivit = []
    for pid in ((m.get('signatures') or {}).get('required_signer_person_ids') or []):
        if pid == pj:
            r = 'puheenjohtaja'
        else:
            r = ROOLINIMET.get(T.rooli_paivana(k, pid, m.get('decision_date')), 'jäsen')
        rivit.append((T.nimi(k, pid), r))
    return rivit


def paatoskappaleet(ty, m, dec, saanto):
    """Yhden asiapykälän rivit: (teksti, lihava). Sanamuotoa ei muuteta, vain jäsennetään."""
    k = ty.get('kokoonpano') or {}
    rivit = [(dec['title'], True)]
    if m.get('decision_method') == 'without_meeting' and dec.get('proposal'):
        rivit.append(('Päätösehdotus: ' + re.sub(r'\s+', ' ', dec['proposal']).strip(), False))
    for osa in jaa_kohdiksi(dec.get('decision_text')):
        rivit.append((osa, False))
    est = T.normalisoi_esteellisyys(dec.get('conflict_of_interest')) or {}
    if est.get('status') == 'declared':
        nimet = ', '.join(T.nimi(k, p) for p in est.get('disqualified_person_ids') or []) or 'Jäsen'
        monikko = len(est.get('disqualified_person_ids') or []) > 1
        t = 'Esteellisyys: %s ilmoitti%s esteellisyydestään eikä osallistu%s asian käsittelyyn' % (
            nimet, 'vat' if monikko else '', 'neet' if monikko else 'nut')
        if est.get('left_room'):
            t += ' ja poistui%s käsittelyn ajaksi' % ('vat' if monikko else '')
        t += '.'
        if est.get('note'):
            t += ' ' + est['note'].strip()
        rivit.append((t, False))
        q = T.johda_paatoksen_paatosvaltaisuus(m, dec, saanto)
        rivit.append(('Asian käsittelyyn osallistui %d hallituksen %d jäsenestä; hallitus oli asiassa päätösvaltainen.'
                      % (q['participating_count'], q['entitled_count']) if q['is_quorate'] else
                      'HUOMIO: asian käsittelyyn osallistui %d hallituksen %d jäsenestä — ei päätösvaltainen.'
                      % (q['participating_count'], q['entitled_count']), False))
    aani = dec.get('vote') or {}
    if aani.get('unanimous') is False:
        t = 'Äänestys: puolesta %s, vastaan %s' % (aani.get('for'), aani.get('against'))
        if aani.get('abstained'):
            t += ', tyhjää %s' % aani['abstained']
        t += '.'
        if aani.get('chair_casting_vote'):
            t += ' Äänten mennessä tasan puheenjohtajan ääni ratkaisi.'
        rivit.append((t, False))
    for e in (aani.get('dissenting_opinions') or []):
        if isinstance(e, dict):
            rivit.append(('Eriävä mielipide (%s): %s' % (T.nimi(k, e.get('person_id')), re.sub(r'\s+', ' ', e.get('text') or '').strip()), False))
    rivit.append(('Päätöstunnus: ' + dec['id'], False))
    return rivit


# ------------------------------------------------------------------ muodostus

def muodosta(ptk, ulos=None, valmis=False, juuri=JUURI, rekisteri=None, tulosta=print):
    """Muodostaa asiakirjan. Palauttaa (polku, löydökset). Nostaa Kieltaytyi, jos --valmis
    ja validointi löysi estäviä virheitä."""
    ty = T.lataa_tyotila(juuri, rekisteri=rekisteri)
    m = T.hae_kokous(ty, ptk)
    if m is None:
        raise SystemExit('Kokousta %s ei löytynyt rekisteristä' % ptk)
    loydokset = T.validoi(ty, kohde=ptk, valmis=valmis)
    virheet = T.virheet(loydokset)
    varoitukset = [l for l in loydokset if l.taso == T.WARNING]
    for l in loydokset:
        if l.taso != T.INFO:
            tulosta(str(l))
    if valmis and virheet:
        tulosta('\nValmista asiakirjaa ei muodostettu: %d estävää virhettä. Korjaa rekisteri ja aja uudelleen.' % len(virheet))
        raise Kieltaytyi(2)

    tapa = m.get('decision_method') if m.get('decision_method') in POHJAT else 'meeting'
    pohja = POHJAT[tapa]
    if not os.path.exists(pohja):
        raise SystemExit('Asiakirjapohja puuttuu: %s (aja python templates/luo-pohja.py)' % pohja)
    ulos = ulos or os.path.join(juuri, ('' if valmis else 'LUONNOS-') + ptk + '.docx')
    shutil.copy(pohja, ulos)
    d = docx.Document(ulos)
    k = ty.get('kokoonpano') or {}
    saanto = T.paatosvaltaisuussaanto(k)
    q = T.johda_kokouksen_paatosvaltaisuus(m, saanto)

    # Perustiedot
    if tapa == 'meeting':
        istunto = m.get('session') or {}
        a, l = T.aikaleima(istunto.get('opened_at')), T.aikaleima(istunto.get('closed_at'))
        if a and l:
            monipaivainen = a.date() != l.date()
            alku = ('%s klo %s' % (T.pvm_fi(a.date()), a.strftime('%H:%M'))) if monipaivainen else ('klo %s' % a.strftime('%H:%M'))
            loppu = ('%s klo %s' % (T.pvm_fi(l.date()), l.strftime('%H:%M'))) if monipaivainen else ('klo %s' % l.strftime('%H:%M'))
            paiva = ('%s – %s' % (T.pvm_fi(a.date()), T.pvm_fi(l.date()))) if monipaivainen else T.pvm_fi(a.date())
        else:
            monipaivainen, alku, loppu, paiva = False, 'klo [puuttuu]', 'klo [puuttuu]', T.pvm_fi(m.get('decision_date'))
        pp = etsi(d, 'Paikka')
        korvaa(pp, '[Kokouspaikka]', m.get('place') or '[kokouspaikka puuttuu]')
        korvaa(etsi(d, 'avasi kokouksen'), 'klo [HH:MM]', alku)
        korvaa(etsi(d, 'päätti kokouksen'), 'klo [HH:MM]', loppu)
        ensimmainen = pp
    else:
        monipaivainen = False
        paiva = T.pvm_fi(m.get('decision_date'))
        ik = m.get('without_meeting') or {}
        ensimmainen = etsi(d, '[DD.MM.YYYY]')
        korvaa(etsi(d, '[LÄHETETTY]'), '[LÄHETETTY]', _aika_fi(ik.get('proposal_sent_at')))
        korvaa(etsi(d, '[MÄÄRÄAIKA]'), '[MÄÄRÄAIKA]', _aika_fi(ik.get('response_deadline')))
        korvaa(etsi(d, '[TODENTAMINEN]'), '[TODENTAMINEN]', ik.get('verification_method') or '[todentamistapa puuttuu]')
    korvaa(etsi(d, '[DD.MM.YYYY]'), '[DD.MM.YYYY]', paiva)
    korvaa(etsi(d, 'numero [#]'), '[#]', juokseva_numero(ptk))

    # Päätösvaltaisuus — johdettu, ei kopioitu
    pv = etsi(d, '[LÄSNÄ]', pakollinen=False) or etsi(d, 'paikalla oli', pakollinen=False)
    if pv is not None:
        if not korvaa(pv, '[LÄSNÄ]', str(q['present_count'])):
            korvaa(pv, '[#]', str(q['present_count']))
        korvaa(pv, '[JÄSENIÄ]', str(q['entitled_count']))
        if not q['is_quorate']:
            pv.add_run(' HUOMIO: LASKENNAN MUKAAN EI PÄÄTÖSVALTAINEN (%d/%d).' % (q['present_count'], q['entitled_count'])).bold = True

    # Osallistujat, poissaolijat, allekirjoittajat — rekisteristä
    rivit, poissa = osallistujarivit(ty, m)
    otsikko = 'Osallistujat' if tapa == 'meeting' else 'Osallistuneet'
    if not taydenna_nimirivit(d, otsikko, rivit):
        tulosta('HUOMIO: pohjassa ei ole "%s"-lohkoa paikkamerkillä %s — osallistujia ei tulostettu' % (otsikko, NIMIPAIKKA))
    pp_ = etsi(d, '[Poissa]', pakollinen=False)
    if pp_ is not None:
        if poissa:
            korvaa(pp_, '[Poissa]', ', '.join(poissa))
        else:
            poista(pp_)
    elif poissa:
        tulosta('HUOMIO: pohjassa ei ole [Poissa]-riviä — poissa olleet (%s) eivät tulostuneet' % ', '.join(poissa))
    if not taydenna_nimirivit(d, 'Allekirjoitukset', allekirjoittajarivit(ty, m), paikka='[Allekirjoittaja]'):
        tulosta('HUOMIO: pohjassa ei ole [Allekirjoittaja]-rivejä — allekirjoittajia ei tulostettu')

    # Asiapykälät
    asiat = [p for p in d.paragraphs if p.text.strip() == '[ASIA #]']
    if not asiat:
        raise SystemExit('Pohjasta ei löytynyt [ASIA #] -kohtia')
    malli = asiat[0]
    ankkuri = malli._p

    def lisaa(teksti, lihava=False, numeroitu=True):
        nonlocal ankkuri
        p = kloonaa_asia(malli, teksti, lihava=lihava, numeroitu=numeroitu)
        ankkuri.addnext(p._p); ankkuri = p._p

    for dec in (m.get('decisions') or []):
        for i, (teksti, lihava) in enumerate(paatoskappaleet(ty, m, dec, saanto)):
            lisaa(teksti, lihava=lihava, numeroitu=(i == 0))
    ilman = m.get('items_without_decision') or []
    if ilman:
        lisaa('Käsitellyt asiat, joista ei tehty päätöstä', lihava=True)
        for osa in ilman:
            lisaa(re.sub(r'\s+', ' ', osa).strip(), numeroitu=False)
    for a in asiat:
        poista(a)

    # Banneri
    if not valmis:
        b = kloonaa_asia(malli, '', numeroitu=False)
        ensimmainen._p.addprevious(b._p)
        b.add_run('LUONNOS — ei allekirjoitettavaksi. Muodostettu päätösrekisteristä (%s). '
                  'Validointi: %d estävää virhettä, %d varoitusta. Tarkistus ja allekirjoitus vahvistavat sanamuodot.'
                  % (ptk, len(virheet), len(varoitukset))).bold = True
        b.paragraph_format.left_indent = Cm(0)

    d.save(ulos)
    tulosta('%s: %s' % ('Asiakirja' if valmis else 'Luonnos', ulos))
    tulosta('  %s %s, %d päätöstä, %s' % ('päätös ilman kokousta' if tapa == 'without_meeting' else 'kokous', ptk,
                                          len(m.get('decisions') or []), 'monipäiväinen' if monipaivainen else 'yksipäiväinen'))
    return ulos, loydokset


def _aika_fi(x):
    a = T.aikaleima(x)
    if a:
        return '%s klo %s' % (T.pvm_fi(a.date()), a.strftime('%H:%M'))
    return T.pvm_fi(x) or '[puuttuu]'


def main(argv=None):
    argv = list(sys.argv[1:] if argv is None else argv)
    valmis = '--valmis' in argv
    rekisteri = None
    if '--rekisteri' in argv:
        i = argv.index('--rekisteri'); rekisteri = argv[i + 1]; argv = argv[:i] + argv[i + 2:]
    argv = [a for a in argv if not a.startswith('--')]
    if not argv:
        raise SystemExit(__doc__)
    try:
        muodosta(argv[0], argv[1] if len(argv) > 1 else None, valmis=valmis, rekisteri=rekisteri)
    except Kieltaytyi:
        return 2
    return 0


if __name__ == '__main__':
    sys.exit(main())
